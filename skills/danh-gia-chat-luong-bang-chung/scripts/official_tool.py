"""Read a lawfully supplied original locally, never download or redistribute it."""
from hashlib import sha256
from io import BytesIO
from pathlib import Path
import re
from zipfile import ZipFile


def load_official_tool(path: Path, expected_sha256: str, *, authorized_use: bool = False) -> dict:
    if authorized_use is not True:
        raise ValueError("LICENSE_SCOPE_REQUIRED")
    if not re.fullmatch(r"[0-9a-fA-F]{64}", expected_sha256):
        raise ValueError("EXPECTED_SOURCE_HASH_REQUIRED")
    path = Path(path)
    if path.stat().st_size > 32 * 1024 * 1024:
        raise ValueError("SOURCE_TOO_LARGE")
    data = path.read_bytes()
    digest = sha256(data).hexdigest()
    if digest != expected_sha256.lower():
        raise ValueError("SOURCE_HASH_MISMATCH")
    report = {"status": "ORIGINAL_LOADED", "sha256": digest, "appraisal_completed": False,
              "redistributed": False, "paragraphs": {}, "tables": {}, "pages": {}}
    if path.suffix.lower() == ".docx":
        with ZipFile(BytesIO(data)) as package:
            if sum(part.file_size for part in package.infolist()) > 64 * 1024 * 1024:
                raise ValueError("SOURCE_PACKAGE_TOO_LARGE")
        from docx import Document
        doc = Document(BytesIO(data))
        report["paragraphs"] = {f"P{i}": p.text for i, p in enumerate(doc.paragraphs)}
        report["tables"] = {f"T{ti}R{ri}C{ci}": cell.text for ti, table in enumerate(doc.tables)
                            for ri, row in enumerate(table.rows) for ci, cell in enumerate(row.cells)}
    elif path.suffix.lower() == ".pdf":
        try:
            import fitz
        except ImportError as exc:
            raise RuntimeError("PDF_READER_UNAVAILABLE") from exc
        with fitz.open(stream=data, filetype="pdf") as doc:
            report["pages"] = {str(i + 1): page.get_text() for i, page in enumerate(doc)}
        if not any(text.strip() for text in report["pages"].values()):
            raise ValueError("SOURCE_TEXT_UNREADABLE_OCR_REQUIRED")
    else:
        raise ValueError("UNSUPPORTED_SOURCE_FORMAT")
    return report
