"""Apply source-traceable, refresh-safe Word mechanics to a new DOCX."""

from copy import deepcopy
from math import sqrt
from pathlib import Path
import sys as _sys
from tempfile import TemporaryDirectory

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor
from docx.text.paragraph import Paragraph

from .structure_profiles import ROOT, _sha256, verify_profile_source, table_border_policy
from .effective_formatting import story_paragraphs, style_chain, clear_font_properties, border_values

_ZOTERO_SCRIPTS = Path(__file__).parents[4] / "skills/quan-ly-trich-dan/scripts"
if str(_ZOTERO_SCRIPTS) not in _sys.path:
    _sys.path.insert(0, str(_ZOTERO_SCRIPTS))

from zotero_fields import audit as _zotero_audit  # noqa: E402

_LEADER_SCRIPTS = Path(__file__).parents[4] / "skills/leader-nghien-cuu/scripts"
if str(_LEADER_SCRIPTS) not in _sys.path:
    _sys.path.insert(0, str(_LEADER_SCRIPTS))
from revision_guard import snapshot, commit_new_artifact  # noqa: E402
from docx_integrity import audit_docx_integrity, compare_docx_integrity


def zotero_guard(source: Path, output: Path) -> dict:
    report = _zotero_audit(source, output)
    if report["status"] == "BLOCKED" and Path(output).resolve() != Path(source).resolve():
        Path(output).unlink(missing_ok=True)
    return report


def _field(paragraph, instruction: str, cached_text: str = ""):
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    paragraph.add_run()._r.append(begin)
    paragraph.add_run()._r.append(instr)
    paragraph.add_run()._r.append(separate)
    if cached_text:
        paragraph.add_run(cached_text)
    paragraph.add_run()._r.append(end)


def _contains_field(element, instruction_prefix: str) -> bool:
    return any(
        (node.text or "").strip().startswith(instruction_prefix)
        for node in element.iter(qn("w:instrText"))
    ) or any(
        (node.get(qn("w:instr")) or "").strip().startswith(instruction_prefix)
        for node in element.iter(qn("w:fldSimple"))
    )


REFRESH_SAFE_STYLES = (
    "Normal",
    "Body Text",
    "Heading 1",
    "Heading 2",
    "Heading 3",
    "Heading 4",
    "TOC Heading",
    "TOC 1",
    "TOC 2",
    "TOC 3",
    "TOC 4",
    "Table of Figures",
    "Caption",
    "Bibliography",
    "List Paragraph",
    "MedRS Table Text",
)


def _set_style_font(style, font_name: str, size_pt: float):
    style.font.name = font_name
    style.font.size = Pt(size_pt)
    style.font.color.rgb = RGBColor(0, 0, 0)
    rpr = style.element.get_or_add_rPr()
    complex_size = rpr.find(qn("w:szCs"))
    if complex_size is None:
        complex_size = OxmlElement("w:szCs")
        rpr.append(complex_size)
    complex_size.set(qn("w:val"), str(int(round(size_pt * 2))))
    rfonts = rpr.get_or_add_rFonts()
    for attribute in ("ascii", "hAnsi", "eastAsia", "cs"):
        rfonts.set(qn(f"w:{attribute}"), font_name)
    for attribute in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme", "csTheme"):
        rfonts.attrib.pop(qn(f"w:{attribute}"), None)


def _ensure_paragraph_style(document, name: str, *, base="Normal"):
    if name in document.styles:
        style = document.styles[name]
    else:
        style = document.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    if base and name != base and base in document.styles:
        style.base_style = document.styles[base]
    return style


def _configure_style_system(document, profile: dict):
    """Make styles, rather than direct run formatting, the single source of truth."""
    body = profile["body"]
    overrides = profile.get("word_styles", {})
    compact_styles = {"Caption", "MedRS Table Text"}
    heading_styles = {f"Heading {level}" for level in range(1, 5)} | {"TOC Heading"}
    gallery_order = {
        "Normal": 1,
        "Body Text": 2,
        "Heading 1": 3,
        "Heading 2": 4,
        "Heading 3": 5,
        "Heading 4": 6,
        "Caption": 7,
        "List Paragraph": 8,
        "MedRS Table Text": 9,
    }
    for name in REFRESH_SAFE_STYLES:
        base = "Normal"
        if name.startswith("TOC ") and name != "TOC Heading":
            base = "Normal"
        style = _ensure_paragraph_style(document, name, base=base)
        setting = overrides.get(name, {})
        _set_style_font(style, setting.get("font", body["font"]), setting.get("size_pt", body["size_pt"]))
        paragraph_format = style.paragraph_format
        paragraph_format.line_spacing = setting.get(
            "line_spacing",
            profile.get("table", {}).get("line_spacing", 1.0) if name in compact_styles else body["line_spacing"],
        )
        paragraph_format.space_before = Pt(setting.get("space_before_pt", 0))
        paragraph_format.space_after = Pt(setting.get("space_after_pt", 0))
        paragraph_format.keep_with_next = name in heading_styles
        style.quick_style = name in {
            "Normal",
            "Body Text",
            "Heading 1",
            "Heading 2",
            "Heading 3",
            "Heading 4",
            "Caption",
            "List Paragraph",
            "MedRS Table Text",
        }
        if name in gallery_order:
            style.hidden = False
            style.unhide_when_used = False
            style.priority = gallery_order[name]

    _normalize_story_fonts(document, profile)


def _normalize_story_fonts(document, profile):
    body = profile["body"]
    paragraphs = list(story_paragraphs(document))
    protected_styles = set()
    field_depth = 0
    for paragraph in paragraphs:
        for run in paragraph.runs:
            boundaries = [n.get(qn("w:fldCharType")) for n in run._r.iter(qn("w:fldChar"))]
            if field_depth > 0 or boundaries or run._r.xpath(".//w:instrText"):
                protected_styles.update(s.style_id for s in style_chain(run.style))
            for boundary in boundaries:
                field_depth += 1 if boundary == "begin" else -1 if boundary == "end" else 0
    field_depth = 0
    for paragraph in paragraphs:
        if paragraph.style and paragraph.style.name not in REFRESH_SAFE_STYLES:
            _set_style_font(paragraph.style, body["font"], body["size_pt"])
        # Reset only managed paragraph spacing. Keep author-specific indentation,
        # pagination, inline emphasis and object-bearing paragraph structure.
        if paragraph.style and paragraph.style.name in REFRESH_SAFE_STYLES:
            paragraph.paragraph_format.line_spacing = None
            paragraph.paragraph_format.space_before = None
            paragraph.paragraph_format.space_after = None
        for run in paragraph.runs:
            boundaries = [node.get(qn("w:fldCharType")) for node in run._r.iter(qn("w:fldChar"))]
            in_field = field_depth > 0 or bool(boundaries)
            for boundary in boundaries:
                if boundary == "begin":
                    field_depth += 1
                elif boundary == "end":
                    field_depth = max(0, field_depth - 1)
            if in_field:
                continue
            # Keep semantic emphasis, but clear font/size drift so the paragraph style governs.
            clear_font_properties(run._r.rPr)
            chain = list(style_chain(run.style))
            for style in chain:
                if style.style_id not in protected_styles:
                    clear_font_properties(style.element.rPr)
            if any(s.style_id in protected_styles and s.element.rPr is not None for s in chain):
                setting = profile.get("word_styles", {}).get(paragraph.style.name, {})
                run.font.name = setting.get("font", body["font"])
                run.font.size = Pt(setting.get("size_pt", body["size_pt"]))
                complex_size = OxmlElement("w:szCs")
                complex_size.set(qn("w:val"), str(int(round(setting.get("size_pt", body["size_pt"]) * 2))))
                run._r.get_or_add_rPr().append(complex_size)
                for slot in ("ascii", "hAnsi", "eastAsia", "cs"):
                    run._r.get_or_add_rPr().get_or_add_rFonts().set(qn("w:" + slot), run.font.name)


def _enable_field_refresh(document):
    settings = document.settings.element
    update = settings.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        settings.append(update)
    update.set(qn("w:val"), "true")


def _next_numbering_id(elements, attribute: str) -> int:
    values = [int(value) for element in elements if (value := element.get(qn(attribute))) is not None]
    return max(values, default=0) + 1


def _create_medrs_bullet_numbering(document, profile: dict) -> int:
    numbering = document.part.numbering_part.element
    existing = numbering.xpath('./w:abstractNum[w:name/@w:val="MedRSBullet"]')
    if existing:
        abstract_id = int(existing[0].get(qn("w:abstractNumId")))
        for candidate in numbering.findall(qn("w:num")):
            reference = candidate.find(qn("w:abstractNumId"))
            if reference is not None and int(reference.get(qn("w:val"))) == abstract_id:
                return int(candidate.get(qn("w:numId")))
    else:
        abstract_id = _next_numbering_id(numbering.findall(qn("w:abstractNum")), "w:abstractNumId")
        abstract = OxmlElement("w:abstractNum")
        abstract.set(qn("w:abstractNumId"), str(abstract_id))
        name = OxmlElement("w:name")
        name.set(qn("w:val"), "MedRSBullet")
        multi = OxmlElement("w:multiLevelType")
        multi.set(qn("w:val"), "multilevel")
        abstract.append(multi)
        abstract.append(name)
        list_settings = profile.get("lists", {})
        levels = int(list_settings.get("bullet_levels", 3))
        base_left = int(list_settings.get("base_left_twips", 720))
        increment = int(list_settings.get("level_increment_twips", 540))
        hanging = int(list_settings.get("hanging_twips", 360))
        for level in range(levels):
            left = base_left + level * increment
            lvl = OxmlElement("w:lvl")
            lvl.set(qn("w:ilvl"), str(level))
            start = OxmlElement("w:start")
            start.set(qn("w:val"), "1")
            num_fmt = OxmlElement("w:numFmt")
            num_fmt.set(qn("w:val"), "bullet")
            lvl_text = OxmlElement("w:lvlText")
            lvl_text.set(qn("w:val"), "•" if level < 2 else "–")
            suff = OxmlElement("w:suff")
            suff.set(qn("w:val"), "tab")
            ppr = OxmlElement("w:pPr")
            tabs = OxmlElement("w:tabs")
            tab = OxmlElement("w:tab")
            tab.set(qn("w:val"), "num")
            tab.set(qn("w:pos"), str(left))
            tabs.append(tab)
            indent = OxmlElement("w:ind")
            indent.set(qn("w:left"), str(left))
            indent.set(qn("w:hanging"), str(hanging))
            ppr.extend([tabs, indent])
            lvl.extend([start, num_fmt, lvl_text, suff, ppr])
            abstract.append(lvl)
        numbering.insert(0, abstract)

    num_id = _next_numbering_id(numbering.findall(qn("w:num")), "w:numId")
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(qn("w:val"), str(abstract_id))
    num.append(abstract_ref)
    numbering.append(num)
    return num_id


def _existing_bullet_num_ids(document) -> set[int]:
    numbering = document.part.numbering_part.element
    bullet_abstract_ids = {
        int(element.get(qn("w:abstractNumId")))
        for element in numbering.findall(qn("w:abstractNum"))
        if any(
            fmt.get(qn("w:val")) == "bullet"
            for fmt in element.iter(qn("w:numFmt"))
        )
    }
    return {
        int(element.get(qn("w:numId")))
        for element in numbering.findall(qn("w:num"))
        if (reference := element.find(qn("w:abstractNumId"))) is not None
        and int(reference.get(qn("w:val"))) in bullet_abstract_ids
    }


def _normalize_bullets(document, profile: dict):
    candidates = []
    bullet_num_ids = _existing_bullet_num_ids(document)
    for paragraph in document.paragraphs:
        style_name = paragraph.style.name if paragraph.style else ""
        num_pr = paragraph._p.pPr.numPr if paragraph._p.pPr is not None else None
        is_numbered_bullet = (
            num_pr is not None and num_pr.numId is not None and int(num_pr.numId.val) in bullet_num_ids
        )
        if style_name.casefold().startswith("list bullet") or is_numbered_bullet:
            suffix = style_name.rsplit(" ", 1)[-1]
            level = int(num_pr.ilvl.val) if is_numbered_bullet and num_pr.ilvl is not None else (
                int(suffix) - 1 if suffix.isdigit() else 0
            )
            candidates.append((paragraph, max(0, min(level, 2))))
    if not candidates:
        return 0
    num_id = _create_medrs_bullet_numbering(document, profile)
    for paragraph, level in candidates:
        paragraph.style = document.styles["List Paragraph"]
        ppr = paragraph._p.get_or_add_pPr()
        for tag in ("w:ind", "w:tabs", "w:numPr"):
            old = ppr.find(qn(tag))
            if old is not None:
                ppr.remove(old)
        num_pr = OxmlElement("w:numPr")
        ilvl = OxmlElement("w:ilvl")
        ilvl.set(qn("w:val"), str(level))
        num = OxmlElement("w:numId")
        num.set(qn("w:val"), str(num_id))
        num_pr.extend([ilvl, num])
        ppr.insert_element_before(num_pr, "w:suppressLineNumbers", "w:pBdr", "w:shd", "w:tabs", "w:spacing", "w:ind", "w:jc", "w:rPr", "w:sectPr")
    return len(candidates)


def _column_weights(table) -> list[float]:
    weights = []
    for column_index in range(len(table.columns)):
        lengths = []
        for row in table.rows:
            text = " ".join(row.cells[column_index].text.split())
            lengths.append(max(1, len(text)))
        maximum = max(lengths, default=1)
        average = sum(lengths) / max(1, len(lengths))
        weights.append(max(1.0, sqrt(maximum) + 0.25 * sqrt(average)))
    return weights


def _bounded_widths(weights: list[float], total_twips: int, *, minimum_share=0.08, maximum_share=0.65) -> list[int]:
    if not weights:
        return []
    # Shares are preferences, not impossible constraints on 20+ columns.
    minimum = total_twips * min(minimum_share, 1 / len(weights))
    maximum = total_twips * max(maximum_share, 1 / len(weights))
    raw = [total_twips * weight / sum(weights) for weight in weights]
    widths = [min(max(value, minimum), maximum) for value in raw]
    for _ in range(4):
        delta = total_twips - sum(widths)
        if abs(delta) < 1:
            break
        eligible = [i for i, value in enumerate(widths) if minimum < value < maximum]
        if not eligible:
            eligible = list(range(len(widths)))
        share = delta / len(eligible)
        for index in eligible:
            widths[index] = min(max(widths[index] + share, minimum), maximum)
    rounded = [int(round(value)) for value in widths]
    rounded[-1] += total_twips - sum(rounded)
    return rounded


def _format_tables(document, profile: dict):
    table_style = document.styles["MedRS Table Text"]
    section = document.sections[-1]
    # Length subtraction returns EMU as an int; 635 EMU equals one twip.
    usable_width = int((section.page_width - section.left_margin - section.right_margin) / 635)
    for table in document.tables:
        if not table.columns:
            continue
        table_settings = profile.get("table", {})
        widths = _bounded_widths(
            _column_weights(table),
            usable_width,
            minimum_share=float(table_settings.get("minimum_column_share", 0.08)),
            maximum_share=float(table_settings.get("maximum_column_share", 0.65)),
        )
        table.autofit = False
        policy = table_border_policy(profile)
        if policy != "AUTHOR_APPROVAL_REQUIRED":
            _set_borders(table._tbl.tblPr, "w:tblBorders", border_values(policy))
        grid_columns = table._tbl.tblGrid.gridCol_lst
        for index, width in enumerate(widths):
            if index < len(grid_columns):
                grid_columns[index].set(qn("w:w"), str(width))
        for row_index, row in enumerate(table.rows):
            grid_index = 0
            seen_cells = set()
            for cell in row.cells:
                cell_key = id(cell._tc)
                if cell_key in seen_cells:
                    continue
                seen_cells.add(cell_key)
                grid_span = cell._tc.get_or_add_tcPr().gridSpan
                span = int(grid_span.val) if grid_span is not None else 1
                cell_twips = sum(widths[grid_index : grid_index + span])
                cell.width = Inches(cell_twips / 1440)
                tcw = cell._tc.get_or_add_tcPr().get_or_add_tcW()
                tcw.set(qn("w:type"), "dxa")
                tcw.set(qn("w:w"), str(cell_twips))
                if policy != "AUTHOR_APPROVAL_REQUIRED":
                    _set_borders(cell._tc.get_or_add_tcPr(), "w:tcBorders",
                                 border_values(policy, row=row_index, rows=len(table.rows)))
                for paragraph in cell.paragraphs:
                    paragraph.style = table_style
                grid_index += span
    _normalize_story_fonts(document, profile)


def _set_borders(properties, tag, values):
    existing = properties.find(qn(tag))
    if existing is not None:
        properties.remove(existing)
    borders = OxmlElement(tag)
    for side, value in values.items():
        edge = OxmlElement("w:" + side)
        edge.set(qn("w:val"), value)
        if value == "single":
            edge.set(qn("w:sz"), "4")
            edge.set(qn("w:color"), "000000")
        borders.append(edge)
    properties.append(borders)


def _add_introduction_section(document, heading: str):
    target = heading.casefold().strip()
    index = next((i for i, p in enumerate(document.paragraphs) if p.text.casefold().strip() == target), None)
    if index is None or index == 0:
        return False
    preceding = document.paragraphs[index - 1]
    ppr = preceding._p.get_or_add_pPr()
    old = ppr.find(qn("w:sectPr"))
    if old is None:
        ppr.append(deepcopy(document.element.body.sectPr))
    body_sectpr = document.element.body.sectPr
    pg_num = body_sectpr.find(qn("w:pgNumType"))
    if pg_num is None:
        pg_num = OxmlElement("w:pgNumType")
        body_sectpr.append(pg_num)
    pg_num.set(qn("w:start"), "1")
    return True


def format_docx(source_path: Path, output_path: Path, profile: dict, *, author_approved: bool = False) -> dict:
    source = Path(source_path).resolve()
    output = Path(output_path).resolve()
    if source == output:
        return {"status": "SOURCE_OVERWRITE_FORBIDDEN"}
    if output.exists():
        return {"status": "OUTPUT_ALREADY_EXISTS"}
    verification = verify_profile_source(ROOT, profile)
    if verification["status"] != "VERIFIED":
        return verification
    if not author_approved:
        return {"status": "AUTHOR_APPROVAL_REQUIRED"}

    source_snapshot = snapshot(source)
    integrity = audit_docx_integrity(source)
    if integrity["status"] != "PASS":
        return {"status": "DOCX_INTEGRITY_FAILED", "integrity": integrity}
    document = Document(source)
    _configure_style_system(document, profile)
    _enable_field_refresh(document)

    page = profile["page"]
    for section in document.sections:
        section.top_margin = Cm(page["top_cm"])
        section.bottom_margin = Cm(page["bottom_cm"])
        section.left_margin = Cm(page["left_cm"])
        section.right_margin = Cm(page["right_cm"])

    bullet_count = _normalize_bullets(document, profile)
    _format_tables(document, profile)

    toc_heading = next((p for p in document.paragraphs if p.text.casefold().strip() == "mục lục"), None)
    if toc_heading is not None and not _contains_field(document.element, "TOC"):
        toc_paragraph = OxmlElement("w:p")
        ppr = toc_paragraph.get_or_add_pPr()
        pstyle = OxmlElement("w:pStyle")
        pstyle.set(qn("w:val"), document.styles["TOC 1"].style_id)
        ppr.append(pstyle)
        toc_heading._p.addnext(toc_paragraph)
        toc_heading.style = document.styles["TOC Heading"]
        _field(Paragraph(toc_paragraph, toc_heading._parent),
               'TOC \\o "1-4" \\h \\z \\u', cached_text=" ")

    if not _add_introduction_section(document, profile["numbering"]["restart_section_heading"]):
        return {"status": "AUTHOR_INPUT_REQUIRED", "reason": "INTRODUCTION_HEADING_NOT_FOUND"}
    sections = document.sections
    body_section = sections[-1]
    body_section.header.is_linked_to_previous = False
    header_paragraph = body_section.header.paragraphs[0]
    header_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if not _contains_field(body_section.header._element, "PAGE"):
        _field(header_paragraph, "PAGE")

    output.parent.mkdir(parents=True, exist_ok=True)
    with TemporaryDirectory(prefix="medrs-format-") as staging:
        staged = Path(staging) / "formatted.docx"
        document.save(staged)
        zotero = zotero_guard(source, staged)
        if zotero["status"] == "BLOCKED":
            return {"status": "ZOTERO_FIELDS_LOST", "zotero": zotero}
        integrity = compare_docx_integrity(source, staged)
        if integrity["status"] != "PASS":
            return {"status": "DOCX_INTEGRITY_FAILED", "integrity": integrity}
        commit_new_artifact(staged, output, source_snapshot, owner="dinh-dang-tai-lieu")
    return {
        "status": "FORMATTED_WITH_UNPERFORMED_CHECKS",
        "backend": "ooxml-only",
        "source_hash": source_snapshot["sha256"],
        "output_hash": _sha256(output),
        "output": str(output),
        "zotero": zotero,
        "style_system": list(REFRESH_SAFE_STYLES),
        "normalized_bullets": bullet_count,
        "adaptive_tables": len(document.tables),
        "unperformed_checks": ["field-update", "pagination", "visual-layout"],
    }
