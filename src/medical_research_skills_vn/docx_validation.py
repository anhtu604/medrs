"""Inspect the actual generated DOCX and rendered evidence."""

from pathlib import Path
from zipfile import ZipFile
from xml.etree import ElementTree

from docx import Document
from docx.oxml.ns import qn

from .structure_profiles import _sha256
from .word_backends import validate_render_claim
from .docx_formatting import REFRESH_SAFE_STYLES, _column_weights


def _near(actual, expected, tolerance=0.03):
    return actual is not None and abs(actual.cm - float(expected)) <= tolerance


def _refresh_safe_styles_pass(document, profile: dict) -> bool:
    body = profile["body"]
    compact = {"Caption", "MedRS Table Text"}
    for name in REFRESH_SAFE_STYLES:
        if name not in document.styles:
            return False
        style = document.styles[name]
        if style.font.name != body["font"]:
            return False
        expected_spacing = profile.get("table", {}).get("line_spacing", 1.0) if name in compact else body["line_spacing"]
        if style.paragraph_format.line_spacing != expected_spacing:
            return False
    return True


def _adaptive_tables_pass(document) -> bool:
    for table in document.tables:
        if len(table.columns) < 2:
            continue
        weights = _column_weights(table)
        grid = table._tbl.tblGrid.gridCol_lst
        if len(grid) != len(weights):
            return False
        widths = [int(item.get(qn("w:w")) or 0) for item in grid]
        if not all(widths):
            return False
        if max(weights) / min(weights) >= 1.8 and max(widths) / min(widths) < 1.3:
            return False
        for row in table.rows:
            if any(paragraph.style.name != "MedRS Table Text" for cell in row.cells for paragraph in cell.paragraphs):
                return False
    return True


def _list_numbering_pass(document, numbering_xml: str) -> bool:
    if "MedRSBullet" not in numbering_xml:
        return True
    if 'w:hanging="360"' not in numbering_xml:
        return False
    root = ElementTree.fromstring(numbering_xml)
    ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    abstract_ids = {
        node.attrib[f"{{{ns['w']}}}abstractNumId"]
        for node in root.findall("w:abstractNum", ns)
        if (name := node.find("w:name", ns)) is not None and name.attrib.get(f"{{{ns['w']}}}val") == "MedRSBullet"
    }
    medrs_num_ids = {
        int(node.attrib[f"{{{ns['w']}}}numId"])
        for node in root.findall("w:num", ns)
        if (reference := node.find("w:abstractNumId", ns)) is not None
        and reference.attrib.get(f"{{{ns['w']}}}val") in abstract_ids
    }
    list_paragraphs = [
        paragraph
        for paragraph in document.paragraphs
        if paragraph._p.pPr is not None
        and paragraph._p.pPr.numPr is not None
        and paragraph._p.pPr.numPr.numId is not None
        and int(paragraph._p.pPr.numPr.numId.val) in medrs_num_ids
    ]
    return bool(list_paragraphs) and all(
        paragraph.style.name == "List Paragraph"
        and paragraph.paragraph_format.left_indent is None
        and paragraph.paragraph_format.first_line_indent is None
        for paragraph in list_paragraphs
    )


def validate_docx(path: Path, profile: dict, *, rendered_artifacts: list, backend: str = "ooxml-only") -> dict:
    path = Path(path)
    document = Document(path)
    normal = document.styles["Normal"]
    body = profile["body"]
    normal_pass = normal.font.name == body["font"] and round(normal.font.size.pt, 1) in profile["body"]["allowed_size_pt"]
    runs = [run for paragraph in document.paragraphs for run in paragraph.runs if run.text]
    direct_formatting_pass = all(
        (run.font.name in {None, body["font"]})
        and (run.font.size is None or round(run.font.size.pt, 1) in body["allowed_size_pt"])
        for run in runs
    )
    page = profile["page"]
    margins_pass = all(
        _near(section.top_margin, page["top_cm"])
        and _near(section.bottom_margin, page["bottom_cm"])
        and _near(section.left_margin, page["left_cm"])
        and _near(section.right_margin, page["right_cm"])
        for section in document.sections
    )
    with ZipFile(path) as archive:
        document_xml = archive.read("word/document.xml").decode("utf-8")
        settings_xml = archive.read("word/settings.xml").decode("utf-8")
        numbering_xml = archive.read("word/numbering.xml").decode("utf-8") if "word/numbering.xml" in archive.namelist() else ""
        header_xml = "".join(
            archive.read(name).decode("utf-8") for name in archive.namelist() if name.startswith("word/header")
        )
    rendered_pdf = next((item for item in rendered_artifacts if str(item).lower().endswith(".pdf")), None)
    rendered_images = [item for item in rendered_artifacts if str(item).lower().endswith((".png", ".jpg", ".jpeg"))]
    render_state = validate_render_claim(backend, rendered_pdf, rendered_images)
    rendered_check = "PASS" if render_state == "RENDERED_ARTIFACTS_AVAILABLE" else render_state
    if render_state == "RENDERED_ARTIFACTS_AVAILABLE":
        pagination_check = "PASS" if backend in {"word-com", "libreoffice-uno"} else "FIELD_UPDATE_UNVERIFIED"
    else:
        pagination_check = render_state
    return {
        "artifact_hash": _sha256(path),
        "mechanical_checks": {
            "normal_style": "PASS" if normal_pass else "FAIL",
            "direct_formatting": "PASS" if direct_formatting_pass else "FAIL",
            "margins": "PASS" if margins_pass else "FAIL",
            "toc_field": "PASS" if "TOC " in document_xml else "FAIL",
            "page_number_field": "PASS" if "PAGE" in header_xml else "FAIL",
            "page_restart_at_introduction": "PASS" if 'w:pgNumType w:start="1"' in document_xml else "FAIL",
            "refresh_safe_styles": "PASS"
            if _refresh_safe_styles_pass(document, profile) and "updateFields" in settings_xml
            else "FAIL",
            "adaptive_table_widths": "PASS" if _adaptive_tables_pass(document) else "FAIL",
            "list_numbering": "PASS" if _list_numbering_pass(document, numbering_xml) else "FAIL",
        },
        "rendered_checks": {"pagination": pagination_check, "visual_layout": rendered_check},
    }
