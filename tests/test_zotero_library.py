import hashlib
import json
import sqlite3
from pathlib import Path

from docx import Document

from zotero_fields import apply_edits, export_paragraphs, zotero_inventory
from zotero_library import open_library, resolve_items


def make_library(path: Path, *, user_id: str | None = "123456") -> Path:
    conn = sqlite3.connect(path)
    conn.executescript(
        """
        CREATE TABLE settings (setting TEXT, key TEXT, value);
        CREATE TABLE libraries (libraryID INTEGER PRIMARY KEY, type TEXT);
        CREATE TABLE "groups" (groupID INTEGER, libraryID INTEGER);
        CREATE TABLE itemTypes (itemTypeID INTEGER PRIMARY KEY, typeName TEXT);
        CREATE TABLE items (itemID INTEGER PRIMARY KEY, itemTypeID INTEGER, libraryID INTEGER, key TEXT);
        CREATE TABLE deletedItems (itemID INTEGER);
        CREATE TABLE fields (fieldID INTEGER PRIMARY KEY, fieldName TEXT);
        CREATE TABLE itemDataValues (valueID INTEGER PRIMARY KEY, value);
        CREATE TABLE itemData (itemID INTEGER, fieldID INTEGER, valueID INTEGER);
        CREATE TABLE creators (creatorID INTEGER PRIMARY KEY, firstName TEXT, lastName TEXT);
        CREATE TABLE creatorTypes (creatorTypeID INTEGER PRIMARY KEY, creatorType TEXT);
        CREATE TABLE itemCreators (itemID INTEGER, creatorID INTEGER, creatorTypeID INTEGER, orderIndex INTEGER);
        INSERT INTO libraries VALUES (1, 'user'), (2, 'group');
        INSERT INTO "groups" VALUES (777, 2);
        INSERT INTO itemTypes VALUES (1, 'journalArticle');
        INSERT INTO items VALUES (15, 1, 1, 'ABCD2345'), (16, 1, 2, 'GRP00001'), (17, 1, 1, 'DEL00001');
        INSERT INTO deletedItems VALUES (17);
        INSERT INTO fields VALUES (1, 'title'), (2, 'date'), (3, 'publicationTitle');
        INSERT INTO itemDataValues VALUES (1, 'Tăng huyết áp ở người cao tuổi'), (2, '2020-05-01 2020-05-01'), (3, 'Tạp chí Nghiên cứu Y học');
        INSERT INTO itemData VALUES (15, 1, 1), (15, 2, 2), (15, 3, 3);
        INSERT INTO creators VALUES (1, 'Văn An', 'Nguyễn');
        INSERT INTO creatorTypes VALUES (1, 'author'), (2, 'editor'), (3, 'translator'), (4, 'contributor');
        INSERT INTO itemCreators VALUES (15, 1, 1, 0);
        """
    )
    conn.execute("INSERT INTO settings VALUES ('account', 'localUserKey', 'LOCALKEY1')")
    if user_id:
        conn.execute("INSERT INTO settings VALUES ('account', 'userID', ?)", (user_id,))
    conn.commit()
    conn.close()
    return path


def test_synced_user_item_resolves_to_its_zotero_uri_and_csl(tmp_path):
    items, missing = resolve_items(make_library(tmp_path / "zotero.sqlite"), ["ABCD2345"])
    assert missing == []
    item = items[0]
    assert item.uri == "http://zotero.org/users/123456/items/ABCD2345"
    assert item.csl["title"] == "Tăng huyết áp ở người cao tuổi"
    assert item.csl["author"] == [{"family": "Nguyễn", "given": "Văn An"}]
    assert item.csl["issued"] == {"date-parts": [[2020]]}
    assert item.csl["type"] == "article-journal"


def test_group_and_unsynced_libraries_use_their_uri_forms(tmp_path):
    items, _ = resolve_items(make_library(tmp_path / "a.sqlite"), ["GRP00001"])
    assert items[0].uri == "http://zotero.org/groups/777/items/GRP00001"
    items, _ = resolve_items(make_library(tmp_path / "b.sqlite", user_id=None), ["ABCD2345"])
    assert items[0].uri == "http://zotero.org/users/local/LOCALKEY1/items/ABCD2345"


def test_deleted_and_unknown_keys_are_reported_not_emitted(tmp_path):
    items, missing = resolve_items(make_library(tmp_path / "z.sqlite"), ["DEL00001", "NOPE0000"])
    assert items == []
    assert missing == ["DEL00001", "NOPE0000"]


def test_library_is_never_modified(tmp_path):
    db = make_library(tmp_path / "zotero.sqlite")
    before = hashlib.sha256(db.read_bytes()).hexdigest()
    resolve_items(db, ["ABCD2345", "GRP00001"])
    assert hashlib.sha256(db.read_bytes()).hexdigest() == before


def test_library_connection_rejects_writes(tmp_path):
    db = make_library(tmp_path / "zotero.sqlite")
    conn = open_library(db)
    try:
        try:
            conn.execute("UPDATE items SET key = 'CHANGED1' WHERE itemID = 15")
        except sqlite3.OperationalError as error:
            assert "readonly" in str(error).lower()
        else:
            raise AssertionError("Zotero library connection allowed a write")
    finally:
        conn.close()


def test_ambiguous_key_is_not_guessed(tmp_path):
    db = make_library(tmp_path / "zotero.sqlite")
    conn = sqlite3.connect(db)
    conn.execute("INSERT INTO items VALUES (18, 1, 2, 'ABCD2345')")
    conn.commit()
    conn.close()
    items, missing = resolve_items(db, ["ABCD2345"])
    assert items == []
    assert missing == ["ABCD2345"]


def test_mixed_creator_roles_preserve_real_author_and_label(tmp_path):
    db = make_library(tmp_path / "zotero.sqlite")
    conn = sqlite3.connect(db)
    conn.executescript(
        """
        INSERT INTO creators VALUES (2, 'Minh', 'Trần'), (3, 'Lan', 'Lê'), (4, 'Bình', 'Phạm');
        UPDATE itemCreators SET orderIndex = 1 WHERE itemID = 15;
        INSERT INTO itemCreators VALUES (15, 2, 2, 0), (15, 3, 3, 2), (15, 4, 4, 3);
        """
    )
    conn.commit()
    conn.close()
    items, missing = resolve_items(db, ["ABCD2345"])
    assert missing == []
    assert items[0].csl["author"] == [{"family": "Nguyễn", "given": "Văn An"}]
    assert items[0].csl["editor"] == [{"family": "Trần", "given": "Minh"}]
    assert items[0].csl["translator"] == [{"family": "Lê", "given": "Lan"}]
    assert "contributor" not in items[0].csl

    source = tmp_path / "source.docx"
    document = Document()
    document.add_paragraph("Một nhận định.")
    document.save(source)
    out = tmp_path / "out.docx"
    apply_edits(source, {0: "Một nhận định ⟦cite:ABCD2345⟧."}, out, resolver=lambda keys: resolve_items(db, keys))
    instruction = next(iter(zotero_inventory(out)))
    payload = json.loads(instruction.split("CSL_CITATION ", 1)[1])
    assert payload["properties"]["formattedCitation"] == "(Nguyễn, 2020)"


def test_unsupported_feed_library_key_is_unresolved(tmp_path):
    db = make_library(tmp_path / "zotero.sqlite")
    conn = sqlite3.connect(db)
    conn.execute("INSERT INTO libraries VALUES (3, 'feed')")
    conn.execute("INSERT INTO items VALUES (18, 1, 3, 'FEED0001')")
    conn.commit()
    conn.close()
    items, missing = resolve_items(db, ["FEED0001"])
    assert items == []
    assert missing == ["FEED0001"]


def test_cite_marker_becomes_a_live_zotero_field(tmp_path):
    db = make_library(tmp_path / "zotero.sqlite")
    source = tmp_path / "draft.docx"
    document = Document()
    document.add_paragraph("Tăng huyết áp phổ biến ở người cao tuổi.")
    document.save(source)

    out = tmp_path / "out.docx"
    report = apply_edits(
        source,
        {0: "Tăng huyết áp phổ biến ở người cao tuổi ⟦cite:ABCD2345⟧. Xem thêm ⟦cite:NOPE0000⟧."},
        out,
        resolver=lambda keys: resolve_items(db, keys),
    )

    assert report["status"] == "PASS"
    assert report["added_count"] == 1
    assert report["unresolved_citations"] == ["NOPE0000"]
    instruction = next(iter(zotero_inventory(out)))
    payload = json.loads(instruction.split("CSL_CITATION ", 1)[1])
    assert payload["citationItems"][0]["uris"] == ["http://zotero.org/users/123456/items/ABCD2345"]
    assert "[CẦN TRÍCH DẪN: NOPE0000]" in export_paragraphs(out)["paragraphs"][0]["text"]
