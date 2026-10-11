import hashlib
import importlib.util
from pathlib import Path
import subprocess
import sys
import site
from zipfile import ZipFile, ZIP_DEFLATED

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

from medical_research_skills_vn.docx_formatting import format_docx
from medical_research_skills_vn.docx_validation import validate_docx
from profile_fixture import load_test_profile as load_structure_profile
from medical_research_skills_vn.packaging import build_cowork_package

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "skills/dinh-dang-tai-lieu/scripts"))


def integrity(path):
    from docx_integrity import audit_docx_integrity
    return audit_docx_integrity(path)


def package(path, edits):
    with ZipFile(path) as source:
        parts = {name: source.read(name) for name in source.namelist()}
    parts.update(edits)
    with ZipFile(path, "w", ZIP_DEFLATED) as output:
        for name, data in parts.items():
            output.writestr(name, data)
    return path


def fixture(path):
    doc = Document()
    doc.add_heading("MỤC LỤC", 1)
    doc.add_heading("ĐẶT VẤN ĐỀ", 1)
    doc.add_paragraph("Synthetic body")
    doc.save(path)
    return path


def test_missing_footnote_target_is_rejected(tmp_path):
    path = fixture(tmp_path / "broken.docx")
    with ZipFile(path) as source:
        xml = source.read("word/document.xml").replace(b"</w:body>", b'<w:p><w:r><w:footnoteReference w:id="9"/></w:r></w:p></w:body>')
    package(path, {"word/document.xml": xml})
    report = integrity(path)
    assert report["status"] == "FAIL"
    assert any(item["code"] == "MISSING_FOOTNOTE_TARGET" for item in report["issues"])


def test_unbalanced_field_in_header_is_rejected(tmp_path):
    path = fixture(tmp_path / "broken.docx")
    package(path, {"word/header1.xml": '<w:hdr xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:p><w:r><w:fldChar w:fldCharType="begin"/><w:instrText> PAGE </w:instrText></w:r></w:p></w:hdr>'})
    report = integrity(path)
    assert report["status"] == "FAIL"
    assert any(item["code"] == "FIELD_UNBALANCED" for item in report["issues"])


def test_missing_comment_target_and_relationship_are_rejected(tmp_path):
    path = fixture(tmp_path / "broken.docx")
    with ZipFile(path) as source:
        xml = source.read("word/document.xml").replace(b"</w:body>", b'<w:p><w:commentRangeStart w:id="4"/><w:r><w:commentReference w:id="4"/></w:r><w:commentRangeEnd w:id="4"/></w:p></w:body>')
    package(path, {"word/document.xml": xml,
                   "word/_rels/custom.xml.rels": '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rIdMissing" Type="fixture" Target="missing.xml"/></Relationships>'})
    report = integrity(path)
    assert {item["code"] for item in report["issues"]} >= {"MISSING_COMMENT_TARGET", "MISSING_RELATIONSHIP_TARGET"}


def test_plain_text_edits_do_not_change_protected_identity(tmp_path):
    first = fixture(tmp_path / "one.docx")
    second = fixture(tmp_path / "two.docx")
    doc = Document(second)
    doc.add_paragraph("Author added text")
    doc.save(second)
    assert integrity(first)["status"] == "PASS"
    assert integrity(first)["protected_fingerprints"] == integrity(second)["protected_fingerprints"]


def test_removed_bookmark_is_detected_even_when_text_matches(tmp_path):
    from docx_integrity import compare_docx_integrity
    source = fixture(tmp_path / "source.docx")
    returned = fixture(tmp_path / "returned.docx")
    with ZipFile(source) as archive:
        xml = archive.read("word/document.xml").replace(b"</w:body>", b'<w:p><w:bookmarkStart w:id="42" w:name="author_anchor"/><w:bookmarkEnd w:id="42"/></w:p></w:body>')
    package(source, {"word/document.xml": xml})
    report = compare_docx_integrity(source, returned)
    assert report["status"] == "FAIL"
    assert "bookmarks" in report["lost_protected_kinds"]


def test_theme_and_direct_spacing_do_not_override_managed_style(tmp_path):
    source = fixture(tmp_path / "source.docx")
    doc = Document(source)
    body = doc.paragraphs[-1]
    body.paragraph_format.line_spacing = 1.0
    body.paragraph_format.space_after = Pt(18)
    run = body.runs[0]
    run.bold = True
    rfonts = run._r.get_or_add_rPr().get_or_add_rFonts()
    rfonts.set(qn("w:asciiTheme"), "minorHAnsi")
    rfonts.set(qn("w:hAnsiTheme"), "minorHAnsi")
    doc.styles["Normal"].element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:asciiTheme"), "majorHAnsi")
    doc.save(source)
    output = tmp_path / "output.docx"
    profile = load_structure_profile(ROOT / "profiles/institution/hmu/word-format-master-2020-current-2026.yaml")
    format_docx(source, output, profile, author_approved=True)
    result = Document(output)
    body = next(p for p in result.paragraphs if p.text == "Synthetic body")
    assert body.paragraph_format.line_spacing is None
    assert body.paragraph_format.space_after is None
    assert body.runs[0].bold is True
    assert body.runs[0]._r.xpath(".//w:rFonts/@w:asciiTheme") == []
    assert result.styles["Normal"].element.xpath(".//w:rFonts/@w:asciiTheme") == []


def test_actual_spacing_drift_after_refresh_fails_validation(tmp_path):
    source = fixture(tmp_path / "source.docx")
    output = tmp_path / "output.docx"
    profile = load_structure_profile(ROOT / "profiles/institution/hmu/word-format-master-2020-current-2026.yaml")
    format_docx(source, output, profile, author_approved=True)
    doc = Document(output)
    doc.paragraphs[-1].paragraph_format.line_spacing = 1.0
    doc.save(output)
    report = validate_docx(output, profile, rendered_artifacts=[])
    assert report["mechanical_checks"]["paragraph_spacing"] == "FAIL"


def test_many_table_columns_are_positive_and_fit_page(tmp_path):
    from medical_research_skills_vn.docx_formatting import _bounded_widths
    widths = _bounded_widths([1] * 20, 9000)
    assert sum(widths) == 9000
    assert min(widths) > 0


def test_cowork_word_runtime_imports_without_root_src(tmp_path):
    archive = build_cowork_package(ROOT, tmp_path / "cowork.zip")
    extracted = tmp_path / "isolated"
    with ZipFile(archive) as bundle:
        bundle.extractall(extracted)
    # -I hides user-installed third-party libraries as well as the repo. Add
    # only the installed dependency directory; never add repo src/PYTHONPATH.
    code = "import sys; sys.path.insert(0,sys.argv[2]); sys.path.insert(0,sys.argv[1]); from medrs_word_runtime.docx_formatting import format_docx; from medrs_word_runtime.docx_validation import validate_docx"
    result = subprocess.run([sys.executable, "-I", "-c", code, str(extracted / "skills/dinh-dang-tai-lieu/scripts"), site.getusersitepackages()], capture_output=True, text=True, cwd=extracted)
    assert result.returncode == 0, result.stderr
