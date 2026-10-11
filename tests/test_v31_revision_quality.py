from datetime import date
from pathlib import Path
import json
import subprocess
import sys
from zipfile import ZipFile, ZIP_DEFLATED

import pytest
from docx import Document

from revision_guard import (snapshot, assert_unchanged, commit_new_artifact,
                            effective_decisions, final_acceptance, receipt_is_current, sha256, output_lease)
from quality_checks import inspect_docx, compare_result_claim, RESULT_ID_FIELDS
from medical_research_skills_vn.routing import RoutingRequest, route_plan, route_request
from medical_research_skills_vn.ethics import validate_source_registers, validate_sources, load_source_register
from medical_research_skills_vn.document_types import resolve_document_type, load_document_type
from medical_research_skills_vn.packaging import build_cowork_package

ROOT = Path(__file__).parents[1]
ACTIVE = {path.parent.name for path in (ROOT / "skills").glob("*/SKILL.md")}


def test_publish_preserves_author_source_and_never_overwrites(tmp_path):
    source, staged, output = (tmp_path / name for name in ("source", "staged", "output"))
    source.write_bytes(b"author edits")
    staged.write_bytes(b"revised")
    before = snapshot(source)
    result = commit_new_artifact(staged, output, before)
    assert result["sha256"] == sha256(staged)
    assert snapshot(source) == before
    with pytest.raises(FileExistsError):
        commit_new_artifact(staged, output, before)
    assert output.read_bytes() == b"revised"
    with pytest.raises(ValueError):
        commit_new_artifact(staged, source, before)


def test_changed_source_blocks_publication(tmp_path):
    source, staged, output = (tmp_path / name for name in ("source", "staged", "output"))
    source.write_bytes(b"old")
    before = snapshot(source)
    source.write_bytes(b"new author edit")
    staged.write_bytes(b"stale edit")
    with pytest.raises(RuntimeError, match="SOURCE_CHANGED"):
        commit_new_artifact(staged, output, before)
    assert not output.exists()
    with pytest.raises(RuntimeError):
        assert_unchanged(before)


def test_concurrent_owner_cannot_claim_same_output(tmp_path):
    output = tmp_path / "result.docx"
    with output_lease(output, "leader"):
        with pytest.raises(FileExistsError):
            with output_lease(output, "other"):
                pytest.fail("Lock was bypassed")
    assert not output.with_name(output.name + ".medrs-lock").exists()


def test_decisions_reuse_latest_confirmed_without_guessing(tmp_path):
    records = [{"scope": "body", "key": "font", "value": 13, "status": "CONFIRMED"},
               {"scope": "table", "key": "font", "value": 12, "status": "CONFIRMED"},
               {"scope": "body", "key": "font", "value": 14, "status": "CONFIRMED"},
               {"scope": "body", "key": "font", "value": 15, "status": "INFERRED"}]
    assert {item["scope"]: item["value"] for item in effective_decisions(records)} == {"body": 14, "table": 12}


def test_acceptance_requires_both_content_and_presentation_on_final_hash(tmp_path):
    artifact = tmp_path / "final.docx"
    artifact.write_bytes(b"final")
    report = tmp_path / "review.md"
    report.write_text("Actual review evidence", encoding="utf-8")
    kinds = ("content", "clear-prose", "scope-preservation", "docx-integrity", "field-refresh", "presentation", "visual-inspection")
    evidence = [{"kind": kind, "artifact_sha256": sha256(artifact), "path": str(report),
                 "sha256": sha256(report), "reviewer": "fixture reviewer", "status": "PASS"} for kind in kinds]
    kwargs = dict(source_preserved=True, propagation_complete=True)
    tasks = [{"id": "T1", "status": "DONE"}]
    receipt = final_acceptance(artifact, tasks, evidence, **kwargs)
    assert receipt["status"] == "ACCEPTED_INTERNAL"
    assert receipt["author_acceptance"] == "PENDING"
    assert receipt_is_current(receipt)
    assert final_acceptance(artifact, tasks, evidence[:-1], **kwargs)["status"] == "VERIFICATION_INCOMPLETE"
    assert final_acceptance(artifact, [{"id": "T1", "status": "OPEN"}], evidence, **kwargs)["status"] == "NEEDS_REVISION"
    artifact.write_bytes(b"repaired after acceptance")
    assert not receipt_is_current(receipt)
    assert final_acceptance(artifact, tasks, evidence, **kwargs)["status"] == "RE_REVIEW_REQUIRED"


def test_workflow_language_split_across_runs_is_detected_but_comments_excluded(tmp_path):
    path = tmp_path / "draft.docx"
    doc = Document()
    paragraph = doc.add_paragraph()
    paragraph.add_run("Ta")
    paragraph.add_run("sk 4")
    doc.save(path)
    with ZipFile(path, "a", ZIP_DEFLATED) as archive:
        archive.writestr("word/comments.xml", '<w:comments xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:p><w:r><w:t>Task 5</w:t></w:r></w:p></w:comments>')
    report = inspect_docx(path)
    assert len(report["findings"]) == 1
    assert report["findings"][0]["text"] == "Task 4"
    assert "visual-layout" in report["unperformed"]


@pytest.mark.parametrize("field", RESULT_ID_FIELDS)
def test_equal_p_from_wrong_result_identity_is_rejected(field):
    actual = {key: "confirmed" for key in RESULT_ID_FIELDS}
    actual["p"] = 0.02
    claim = dict(actual, **{field: "different"})
    assert "IDENTITY_MISMATCH:" + field in compare_result_claim(claim, actual)["errors"]


def test_proportion_ci_is_flagged_not_clipped_and_change_score_not_misbounded():
    actual = {key: "confirmed" for key in RESULT_ID_FIELDS}
    actual.update(ci_low=90, ci_high=110, estimand_bounds=[0, 100])
    assert "CI_OUT_OF_DOMAIN" in compare_result_claim(actual, actual)["errors"]
    assert actual["ci_high"] == 110
    actual.pop("estimand_bounds")
    assert compare_result_claim(actual, actual)["status"] == "PASS"


@pytest.mark.parametrize("invalid", [float("nan"), float("inf"), "0.02", None, True])
def test_nonfinite_and_non_numeric_claim_values_fail_without_crashing(invalid):
    actual = {key: "confirmed" for key in RESULT_ID_FIELDS}
    actual.update(p=0.02)
    claim = dict(actual, p=invalid)
    assert "NUMERIC_INVALID:p" in compare_result_claim(claim, actual)["errors"]


def test_compound_plan_keeps_all_actions_and_does_not_split_conclusion_nouns():
    plan = route_plan(RoutingRequest("viết ppnc; căn bảng; nghiệm thu file cuối"), {}, ACTIVE)
    assert [task["skill"] for task in plan["tasks"]] == ["viet-phuong-phap", "dinh-dang-tai-lieu", "leader-nghien-cuu"]
    assert plan["coordinator"] == "leader-nghien-cuu"
    assert len(route_plan(RoutingRequest("viết kết luận và khuyến nghị"), {}, ACTIVE)["tasks"]) == 1


def test_explicit_route_beats_incidental_keywords_and_exact_legacy_boundary():
    assert route_request(RoutingRequest("/viet-phuong-phap có bảng trong Word"), {}, ACTIVE).canonical == "viet-phuong-phap"
    legacy = {"old": {"canonical": "viet-phuong-phap", "mode": "legacy"}}
    assert route_request(RoutingRequest("older request"), legacy, ACTIVE).canonical == "medrs"


@pytest.mark.parametrize("text,expected", [
    ("tìm tài liệu cho phần phương pháp", "tim-y-van"),
    ("sửa bàn luận theo STROBE", "viet-ban-luan"),
    ("rút gọn tổng quan y văn", "viet-tong-quan"),
    ("sửa ppnc có bảng", "viet-phuong-phap"),
    ("revise the discussion using CONSORT", "viet-ban-luan"),
])
def test_action_owner_beats_incidental_section_or_checklist_mentions(text, expected):
    assert route_request(RoutingRequest(text), {}, ACTIVE).canonical == expected


def test_quarantined_source_can_ship_but_runtime_keeps_stale_warning():
    register = load_source_register(ROOT / "skills/dao-duc-va-quan-tri-du-lieu/references/source-register.yaml")
    report = validate_sources(register, date(2026, 10, 11))
    assert report.has("SOURCE_STALE")
    assert "SOURCE_STALE: vn-tt43-2024-byt" in report.banners
    assert validate_source_registers(ROOT, date(2026, 10, 11)) == []


@pytest.mark.parametrize("document_label,expected", [("đề tài CK2", "project-ck2"),
                                             ("đề tài cơ sở", "project-institutional"),
                                             ("luận văn CK2", "thesis-specialist")])
def test_basic_projects_do_not_replace_official_specialist_thesis(document_label, expected):
    resolved = resolve_document_type(document_label, "vi-medical-academic")
    assert resolved == expected
    profile = load_document_type(ROOT, resolved)
    if resolved.startswith("project-"):
        assert profile["convention"]["complexity"] == "basic-question-led"
        assert profile["convention"]["contributions_section"] is None
        assert profile["rules"] is None


def test_cowork_zip_router_executes_without_source_tree(tmp_path):
    bundle = build_cowork_package(ROOT, tmp_path / "cowork.zip")
    extracted = tmp_path / "sandbox"
    with ZipFile(bundle) as archive:
        archive.extractall(extracted)
    assert not (extracted / "src").exists()
    run = subprocess.run([sys.executable, str(extracted / "skills/medrs/scripts/router.py"),
                          "viết ppnc; căn bảng; nghiệm thu file cuối"],
                         cwd=extracted, capture_output=True, text=True, encoding="utf-8")
    assert run.returncode == 0, run.stderr
    plan = json.loads(run.stdout)
    assert [task["skill"] for task in plan["tasks"]] == ["viet-phuong-phap", "dinh-dang-tai-lieu", "leader-nghien-cuu"]
    assert (extracted / "skills/leader-nghien-cuu/scripts/revision_guard.py").is_file()
    assert (extracted / "profiles/document-type/project-ck2.yaml").is_file()

