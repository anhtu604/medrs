from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


def handoff_api():
    import external_revision
    return external_revision


def consent():
    return {"confirmed": True, "destination": "author-selected-ai", "data_class": "synthetic", "fields": ["manuscript"], "scope": "Revise prose only"}


def text_file(path, text):
    path.write_text(text, encoding="utf-8")
    return path


def test_no_transfer_without_specific_consent(tmp_path):
    api = handoff_api()
    source = text_file(tmp_path / "draft.txt", "Synthetic draft")
    for auth in ({"confirmed": False}, {"confirmed": True}, {**consent(), "fields": []}):
        report = api.prepare_handoff(source, [], auth)
        assert report["status"] == "AUTHOR_APPROVAL_REQUIRED"
        assert report["transfer_permitted"] is False


def test_handoff_binds_source_and_tasks_but_does_not_send(tmp_path):
    api = handoff_api()
    source = text_file(tmp_path / "draft.txt", "Synthetic draft")
    report = api.prepare_handoff(source, [{"id": "style", "scope": "prose"}], consent())
    assert report["status"] == "READY_FOR_AUTHORIZED_HANDOFF"
    assert report["sent"] is False
    assert report["source_sha256"]
    assert report["task_ids"] == ["style"]
    assert "Synthetic draft" not in report["prompt"]


def test_changed_input_not_same_as_last_delivered_source(tmp_path):
    api = handoff_api()
    source = text_file(tmp_path / "draft.txt", "Initial")
    handoff = api.prepare_handoff(source, [], consent())
    source.write_text("Author changed input", encoding="utf-8")
    returned = text_file(tmp_path / "returned.txt", "Returned edit")
    report = api.inspect_returned_revision(source, returned, handoff)
    assert report["status"] == "SOURCE_CHANGED"
    assert report["acceptance_invalidated"] is True


def test_numeric_edits_require_content_review(tmp_path):
    api = handoff_api()
    source = text_file(tmp_path / "draft.txt", "Available N = 20")
    returned = text_file(tmp_path / "returned.txt", "Available N = 30")
    report = api.inspect_returned_revision(source, returned, api.prepare_handoff(source, [{"id": "prose"}], consent()))
    assert report["status"] == "REVIEW_REQUIRED"
    assert report["numeric_change"] is True
    assert report["acceptance_invalidated"] is True
    assert report["task_ids"] == ["prose"]


def test_plain_prose_edit_is_not_automatically_accepted(tmp_path):
    api = handoff_api()
    source = text_file(tmp_path / "draft.txt", "Initial wording")
    returned = text_file(tmp_path / "returned.txt", "Author revised wording")
    report = api.inspect_returned_revision(source, returned, api.prepare_handoff(source, [], consent()))
    assert report["status"] == "REVIEW_REQUIRED"
    assert report["diff"]
    assert "scope-and-content" in report["unperformed_checks"]


def field_doc(path, key):
    doc = Document()
    paragraph = doc.add_paragraph("Synthetic source ")
    for tag, attr, value, text in [("fldChar", "fldCharType", "begin", None), ("instrText", None, None, 'ADDIN ZOTERO_ITEM CSL_CITATION {"citationItems":[{"id":' + str(key) + '}]}'), ("fldChar", "fldCharType", "separate", None), ("t", None, None, "[1]"), ("fldChar", "fldCharType", "end", None)]:
        node = OxmlElement("w:" + tag)
        if attr:
            node.set(qn("w:" + attr), value)
        node.text = text
        paragraph.add_run()._r.append(node)
    doc.save(path)
    return path


def test_same_display_number_different_citation_identity_is_blocked(tmp_path):
    api = handoff_api()
    source = field_doc(tmp_path / "source.docx", 15)
    returned = field_doc(tmp_path / "returned.docx", 16)
    report = api.inspect_returned_revision(source, returned, api.prepare_handoff(source, [], consent()))
    assert report["status"] == "PROTECTED_CONTENT_CHANGED"
    assert "fields" in report["protected_changes"]


def test_removed_live_field_is_blocked(tmp_path):
    api = handoff_api()
    source = field_doc(tmp_path / "source.docx", 15)
    returned = tmp_path / "returned.docx"
    doc = Document()
    doc.add_paragraph("Synthetic source [1]")
    doc.save(returned)
    report = api.inspect_returned_revision(source, returned, api.prepare_handoff(source, [], consent()))
    assert report["status"] == "PROTECTED_CONTENT_CHANGED"


def test_removed_field_inside_track_changes_is_blocked(tmp_path):
    api = handoff_api()
    source = field_doc(tmp_path / "source.docx", 15)
    doc = Document(source)
    paragraph = doc.paragraphs[0]
    ins = OxmlElement("w:ins")
    ins.set(qn("w:id"), "7")
    ins.set(qn("w:author"), "Synthetic reviewer")
    for run in list(paragraph._p)[1:]:
        ins.append(run)
    paragraph._p.append(ins)
    doc.save(source)
    returned = tmp_path / "returned.docx"
    clean = Document()
    clean.add_paragraph("Synthetic source [1]")
    clean.save(returned)
    report = api.inspect_returned_revision(source, returned, api.prepare_handoff(source, [], consent()))
    assert report["status"] == "PROTECTED_CONTENT_CHANGED"


def test_comment_deleted_from_returned_file_is_blocked(tmp_path):
    api = handoff_api()
    source = tmp_path / "source.docx"
    doc = Document()
    doc.add_paragraph("Synthetic text")
    doc.save(source)
    returned = tmp_path / "returned.docx"
    returned.write_bytes(source.read_bytes())
    with ZipFile(source) as archive:
        parts = {name: archive.read(name) for name in archive.namelist()}
    parts["word/document.xml"] = parts["word/document.xml"].replace(b"</w:body>", b'<w:p><w:commentRangeStart w:id="4"/><w:r><w:commentReference w:id="4"/></w:r><w:commentRangeEnd w:id="4"/></w:p></w:body>')
    parts["word/comments.xml"] = b'<w:comments xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:comment w:id="4"><w:p><w:r><w:t>Synthetic comment</w:t></w:r></w:p></w:comment></w:comments>'
    with ZipFile(source, "w", ZIP_DEFLATED) as archive:
        for name, content in parts.items():
            archive.writestr(name, content)
    report = api.inspect_returned_revision(source, returned, api.prepare_handoff(source, [], consent()))
    assert report["status"] == "PROTECTED_CONTENT_CHANGED"
    assert "comments" in report["protected_changes"]
