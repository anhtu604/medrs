from pathlib import Path
import hashlib
import sys
import pytest
from docx import Document

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "skills/danh-gia-chat-luong-bang-chung/scripts"))


def test_matching_original_is_actually_loaded_with_locators(tmp_path):
    from official_tool import load_official_tool
    path = tmp_path / "synthetic.docx"
    doc = Document()
    doc.add_paragraph("Synthetic official-tool fixture, not real guidance")
    doc.add_table(rows=1, cols=1).cell(0, 0).text = "Fixture question"
    doc.save(path)
    digest = hashlib.sha256(path.read_bytes()).hexdigest()
    report = load_official_tool(path, digest, authorized_use=True)
    assert report["status"] == "ORIGINAL_LOADED"
    assert report["sha256"] == digest
    assert report["paragraphs"]["P0"].startswith("Synthetic")
    assert report["tables"]["T0R0C0"] == "Fixture question"
    assert report["appraisal_completed"] is False


def test_mismatched_original_rejected(tmp_path):
    from official_tool import load_official_tool
    path = tmp_path / "wrong.docx"
    path.write_bytes(b"not a matching original")
    with pytest.raises(ValueError, match="SOURCE_HASH_MISMATCH"):
        load_official_tool(path, "a" * 64, authorized_use=True)


def test_use_authority_required_before_reading_original(tmp_path):
    from official_tool import load_official_tool
    with pytest.raises(ValueError, match="LICENSE_SCOPE_REQUIRED"):
        load_official_tool(tmp_path / "missing.docx", "a" * 64, authorized_use=False)
