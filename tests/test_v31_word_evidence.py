from pathlib import Path
import pytest
from docx import Document
from medical_research_skills_vn.docx_formatting import format_docx
from medical_research_skills_vn.docx_validation import validate_docx
from medical_research_skills_vn.structure_profiles import load_structure_profile, _sha256

ROOT = Path(__file__).parents[1]


def prepare(tmp_path):
    source, output = tmp_path / "source.docx", tmp_path / "output.docx"
    doc = Document()
    doc.add_heading("MỤC LỤC", 1)
    doc.add_paragraph("Contents")
    doc.add_heading("ĐẶT VẤN ĐỀ", 1)
    doc.add_paragraph("Author content")
    doc.save(source)
    profile = load_structure_profile(ROOT / "profiles/institution/hmu/word-format-master-2020-current-2026.yaml")
    return source, output, profile


def test_existing_word_output_is_not_replaced(tmp_path):
    source, output, profile = prepare(tmp_path)
    output.write_bytes(b"existing author version")
    before = output.read_bytes()
    assert format_docx(source, output, profile, author_approved=True)["status"] == "OUTPUT_ALREADY_EXISTS"
    assert output.read_bytes() == before


def test_refresh_safe_style_validation_respects_role_specific_overrides(tmp_path):
    source, output, profile = prepare(tmp_path)
    profile["word_styles"] = {"Caption": {"font": "Arial", "size_pt": 12, "line_spacing": 1.0},
                              "TOC 1": {"size_pt": 13, "line_spacing": 1.5}}
    format_docx(source, output, profile, author_approved=True)
    report = validate_docx(output, profile, rendered_artifacts=[])
    assert report["mechanical_checks"]["refresh_safe_styles"] == "PASS"
    doc = Document(output)
    assert doc.styles["Caption"].font.name == "Arial"
    assert doc.styles["TOC 1"].font.size.pt == 13


@pytest.mark.parametrize("backend", ["word-com", "libreoffice-uno"])
def test_backend_name_and_render_files_alone_do_not_prove_review(tmp_path, backend):
    source, output, profile = prepare(tmp_path)
    format_docx(source, output, profile, author_approved=True)
    pdf, png = tmp_path / "render.pdf", tmp_path / "page.png"
    pdf.write_bytes(b"%PDF-1.4 fixture")
    png.write_bytes(b"fixture image")
    evidence_path = tmp_path / "review.md"
    evidence_path.write_text("Fixture evidence; no live Word render claimed", encoding="utf-8")
    artifacts = [pdf, png]
    kwargs = dict(rendered_artifacts=artifacts, backend=backend)
    checks = validate_docx(output, profile, **kwargs)["rendered_checks"]
    assert checks == {"pagination": "FIELD_UPDATE_UNVERIFIED", "visual_layout": "VISUAL_INSPECTION_UNPERFORMED"}
    evidence = {"status": "PASS", "reviewer": "fixture reviewer", "artifact_sha256": _sha256(output),
                "rendered_sha256": {str(path.resolve()): _sha256(path) for path in artifacts},
                "path": str(evidence_path), "sha256": _sha256(evidence_path)}
    checked = validate_docx(output, profile, field_update_evidence=evidence, visual_review=evidence, **kwargs)
    assert checked["rendered_checks"] == {"pagination": "PASS", "visual_layout": "PASS"}
    png.write_bytes(b"changed render")
    assert validate_docx(output, profile, field_update_evidence=evidence, visual_review=evidence, **kwargs)["rendered_checks"] == checks
