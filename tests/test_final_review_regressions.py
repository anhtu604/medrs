"""Real synthetic packages reproduce the six independent review findings."""
from copy import deepcopy
from pathlib import Path
import runpy
import sys
from zipfile import ZipFile

import pytest
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

from profile_fixture import load_test_profile
from test_docx_revision_traps import fixture, package
from medical_research_skills_vn.docx_formatting import format_docx
from medical_research_skills_vn.docx_validation import validate_docx
from medical_research_skills_vn.structure_profiles import apply_docx_structure, verify_profile_source

ROOT = Path(__file__).resolve().parents[1]
for skill in ("ho-so-nghien-cuu", "dinh-dang-tai-lieu"):
    sys.path.insert(0, str(ROOT / "skills" / skill / "scripts"))


def profile():
    return load_test_profile(ROOT / "profiles/institution/hmu/word-format-master-2020-current-2026.yaml")


def decision():
    return {"confirmed": True, "data_updated_to_real": True, "author": "synthetic", "timestamp": "2026-10-11"}


def test_repeated_declaration_cannot_clear_simulation_conflict():
    from data_provenance import transition_data_state
    state = {"state": "SIMULATED", "digest": "simulated-a"}
    replacement = {"digest": "simulated-a", "replacement_received": True}
    for _ in range(3):
        state = transition_data_state(state, replacement, decision())
        assert state["state"] == "UNRESOLVED"
        assert not state["remove_current_simulation_labels"]


def test_historical_simulated_digest_requires_scoped_resolution():
    from data_provenance import transition_data_state
    state = transition_data_state({"state": "SIMULATED", "digest": "a"},
                                  {"digest": "b", "replacement_received": True}, decision())
    returned = transition_data_state(state, {"digest": "a", "replacement_received": True}, decision())
    assert returned["state"] == "UNRESOLVED"
    assert not returned["remove_current_simulation_labels"]
    assert "SAME_VALUES_REQUIRE_REVIEW" in returned["qualifications"]


@pytest.mark.parametrize("change,expected", [({}, "AUTHOR_DECLARED_REAL"), ({"digest": "other"}, "UNRESOLVED"),
                                            ({"scope": "paragraph"}, "UNRESOLVED"), ({"reason": ""}, "UNRESOLVED")])
def test_known_conflict_resolution_is_scoped_bound_and_retained(change, expected):
    from data_provenance import transition_data_state
    previous = {"state": "SIMULATED", "digest": "a"}
    replacement = {"digest": "a", "replacement_received": True}
    previous = transition_data_state(previous, replacement, decision())
    before = deepcopy(previous)
    resolution = {"status": "CONFIRMED", "scope": "dataset-replacement", "digest": "a",
                  "author": "synthetic", "timestamp": "2026-10-11", "reason": "Synthetic author resolution"} | change
    report = transition_data_state(previous, replacement, decision() | {"conflict_resolution": resolution})
    assert report["state"] == expected
    assert previous == before
    assert report["remove_current_simulation_labels"] == (expected == "AUTHOR_DECLARED_REAL")
    if expected == "AUTHOR_DECLARED_REAL":
        assert report["history"][-1]["conflict_resolution"] == resolution


def test_validator_resolves_east_asian_theme_not_latin_face(tmp_path):
    from lxml import etree
    source = character_fixture(tmp_path / "source.docx")
    output = tmp_path / "output.docx"
    rules = profile()
    format_docx(source, output, rules, author_approved=True)
    doc = Document(output)
    style = doc.styles["Synthetic character drift"]
    rf = style.element.get_or_add_rPr().get_or_add_rFonts()
    rf.set(qn("w:eastAsiaTheme"), "majorEastAsia")
    doc.save(output)
    with ZipFile(output) as archive:
        theme = etree.fromstring(archive.read("word/theme/theme1.xml"))
    ns = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
    theme.find(".//a:majorFont/a:latin", ns).set("typeface", rules["body"]["font"])
    theme.find(".//a:majorFont/a:ea", ns).set("typeface", "Arial")
    package(output, {"word/theme/theme1.xml": etree.tostring(theme)})
    assert validate_docx(output, rules, rendered_artifacts=[])["mechanical_checks"]["direct_formatting"] == "FAIL"


@pytest.mark.parametrize("kind", ["footnote", "endnote"])
def test_removed_note_reference_with_definition_retained_fails(tmp_path, kind):
    from docx_integrity import compare_docx_integrity
    source = fixture(tmp_path / "source.docx")
    with ZipFile(source) as archive:
        xml = archive.read("word/document.xml")
    reference = f'<w:p><w:r><w:{kind}Reference w:id="9"/></w:r></w:p>'.encode()
    definition = f'<w:{kind}s xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:{kind} w:id="9"><w:p><w:r><w:t>Synthetic note</w:t></w:r></w:p></w:{kind}></w:{kind}s>'
    package(source, {"word/document.xml": xml.replace(b"</w:body>", reference + b"</w:body>"),
                     f"word/{kind}s.xml": definition})
    output = tmp_path / "output.docx"
    output.write_bytes(source.read_bytes())
    package(output, {"word/document.xml": xml})
    report = compare_docx_integrity(source, output)
    assert report["status"] == "FAIL"
    assert kind + "s" in report["lost_protected_kinds"]


@pytest.mark.parametrize("mutate", [
    {"source_kind": "training-regulation"},
    {"document_type": "ckii-thesis"},
    {"semantic_rules": []},
])
def test_executable_source_gate_rejects_inapplicable_or_empty_rules(tmp_path, mutate):
    rules = profile() | mutate
    assert verify_profile_source(ROOT, rules)["status"] == "OFFICIAL_RULE_REQUIRED"
    source = fixture(tmp_path / "source.docx")
    for operation in (format_docx, apply_docx_structure):
        output = tmp_path / (operation.__name__ + ".docx")
        assert operation(source, output, rules, author_approved=True)["status"] == "OFFICIAL_RULE_REQUIRED"
        assert not output.exists()


def split_field(path, first, second, *, end=True):
    source = fixture(path)
    with ZipFile(source) as archive:
        xml = archive.read("word/document.xml")
    field = ('<w:p><w:r><w:fldChar w:fldCharType="begin"/></w:r>'
             f'<w:r><w:instrText>{first}</w:instrText></w:r>'
             f'<w:r><w:instrText>{second}</w:instrText></w:r>'
             + ('<w:r><w:fldChar w:fldCharType="end"/></w:r>' if end else '') + '</w:p>').encode()
    return package(source, {"word/document.xml": xml.replace(b"</w:body>", field + b"</w:body>")})


def test_uno_blocks_split_external_instruction(tmp_path):
    guard = runpy.run_path(str(ROOT / "skills/dinh-dang-tai-lieu/scripts/refresh_uno.py"))["_safe_source"]
    with pytest.raises(ValueError, match="EXTERNAL_UPDATE"):
        guard(split_field(tmp_path / "external.docx", "D", "DEAUTO synthetic"))


def test_uno_detects_split_toc(tmp_path):
    guard = runpy.run_path(str(ROOT / "skills/dinh-dang-tai-lieu/scripts/refresh_uno.py"))["_safe_source"]
    assert guard(split_field(tmp_path / "toc.docx", "T", 'OC \\o "1-4"')) is True


def test_uno_refuses_unbalanced_instruction(tmp_path):
    guard = runpy.run_path(str(ROOT / "skills/dinh-dang-tai-lieu/scripts/refresh_uno.py"))["_safe_source"]
    with pytest.raises(ValueError, match="FIELD_UNBALANCED"):
        guard(split_field(tmp_path / "malformed.docx", "T", "OC", end=False))


def character_fixture(path):
    source = fixture(path)
    doc = Document(source)
    style = doc.styles.add_style("Synthetic character drift", WD_STYLE_TYPE.CHARACTER)
    style.font.name, style.font.size, style.font.italic = "Arial", Pt(28), True
    doc.paragraphs[-1].runs[0].style = style
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "A", "Synthetic longer text"
    table.cell(0, 1).paragraphs[0].runs[0].style = style
    doc.sections[0].header.paragraphs[0].add_run("Synthetic header").style = style
    doc.save(source)
    return source


def test_formatter_normalizes_character_style_fonts_not_emphasis(tmp_path):
    source = character_fixture(tmp_path / "source.docx")
    output = tmp_path / "output.docx"
    rules = profile()
    assert format_docx(source, output, rules, author_approved=True)["status"] == "FORMATTED_WITH_UNPERFORMED_CHECKS"
    doc = Document(output)
    style = doc.styles["Synthetic character drift"]
    assert style.font.name in {None, rules["body"]["font"]}
    assert style.font.size is None or style.font.size.pt == rules["body"]["size_pt"]
    assert style.font.italic is True
    assert validate_docx(output, rules, rendered_artifacts=[])["mechanical_checks"]["direct_formatting"] == "PASS"


def test_validator_detects_character_style_drift_in_table(tmp_path):
    source = character_fixture(tmp_path / "source.docx")
    output = tmp_path / "output.docx"
    rules = profile()
    format_docx(source, output, rules, author_approved=True)
    assert validate_docx(output, rules, rendered_artifacts=[])["mechanical_checks"]["direct_formatting"] == "PASS"
    doc = Document(output)
    style = doc.styles.add_style("Returned table drift", WD_STYLE_TYPE.CHARACTER)
    style.font.name, style.font.size = "Arial", Pt(28)
    doc.tables[0].cell(0, 1).paragraphs[0].runs[0].style = style
    doc.save(output)
    assert validate_docx(output, rules, rendered_artifacts=[])["mechanical_checks"]["direct_formatting"] == "FAIL"


@pytest.mark.parametrize("policy", ["grid", "no-vertical", "minimal"])
def test_selected_table_border_policy_applied_to_actual_cells(tmp_path, policy):
    from medical_research_skills_vn.docx_formatting import _format_tables
    source = fixture(tmp_path / "source.docx")
    doc = Document(source)
    table = doc.add_table(rows=2, cols=2)
    table.style = "Table Grid"
    for row in table.rows:
        for cell in row.cells:
            cell.text = "Synthetic"
    doc.save(source)
    output = tmp_path / "output.docx"
    rules = profile() | {"table": {**profile()["table"], "table_borders": policy}}
    format_docx(source, output, rules, author_approved=True)
    actual = Document(output).tables[0]
    borders = actual._tbl.tblPr.find(qn("w:tblBorders"))
    assert borders is not None
    for side in ("left", "right", "insideV"):
        assert borders.find(qn("w:" + side)).get(qn("w:val")) == ("single" if policy == "grid" else "nil")
    for row in actual.rows:
        for cell in row.cells:
            assert cell._tc.tcPr.find(qn("w:tcBorders")) is not None
    report = validate_docx(output, rules, rendered_artifacts=[])
    assert report["mechanical_checks"]["table_borders"] == "PASS"


def test_validator_rejects_returned_vertical_cell_border(tmp_path):
    source = character_fixture(tmp_path / "source.docx")
    output = tmp_path / "output.docx"
    rules = profile() | {"table_borders": "no-vertical"}
    format_docx(source, output, rules, author_approved=True)
    doc = Document(output)
    pr = doc.tables[0].cell(0, 0)._tc.get_or_add_tcPr()
    borders = pr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        pr.append(borders)
    left = borders.find(qn("w:left"))
    if left is None:
        left = OxmlElement("w:left")
        borders.append(left)
    left.set(qn("w:val"), "single")
    doc.save(output)
    assert validate_docx(output, rules, rendered_artifacts=[])["mechanical_checks"]["table_borders"] == "FAIL"
