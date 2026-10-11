"""Source evidence is required before claiming official layout verification."""
from hashlib import sha256
from pathlib import Path
import pytest

from docx import Document

from medical_research_skills_vn.structure_profiles import verify_profile_source, load_structure_profile, apply_docx_structure
from medical_research_skills_vn.coverage import load_coverage_manifest, validate_coverage_manifest

ROOT = Path(__file__).parents[1]


def test_registered_metadata_without_source_cannot_certify_layout(tmp_path):
    profile = {
        "verification_status": "CURRENT", "last_verified": "2026-10-11", "expires_after_days": 180,
        "source_path": "missing.docx", "source_url": "https://example.invalid/rules",
        "source_sha256": "A" * 64,
    }
    result = verify_profile_source(tmp_path, profile)
    assert result["status"] == "OFFICIAL_RULE_REQUIRED"
    assert result["reason"] == "SOURCE_MISSING"
    assert result["source_available"] is False


def test_available_original_requires_hash_and_valid_locator(tmp_path):
    source = tmp_path / "synthetic-rules.docx"
    document = Document()
    document.add_paragraph("Synthetic fixture: Methods precedes Results.")
    document.save(source)
    profile = {
        "verification_status": "CURRENT", "last_verified": "2026-10-11", "expires_after_days": 180,
        "source_path": source.name, "source_sha256": sha256(source.read_bytes()).hexdigest(),
        "semantic_rules": [{"id": "order", "source_locator": "P0", "summary": "fixture order"}],
    }
    assert verify_profile_source(tmp_path, profile)["status"] == "VERIFIED"
    profile["semantic_rules"][0]["source_locator"] = "P9"
    assert verify_profile_source(tmp_path, profile)["reason"] == "SOURCE_LOCATOR_INVALID"


def test_vjpm_original_moves_actual_blocks_and_preserves_table(tmp_path):
    profile = load_structure_profile(ROOT / "profiles/journal/vjpm-original-2026.yaml")
    # Explicit synthetic source: this test proves moves, not journal-source authenticity.
    rules = tmp_path / "synthetic-rules.html"
    rules.write_text("<html><body>Synthetic order</body></html>", encoding="utf-8")
    profile.update(source_path=str(rules), source_sha256=sha256(rules.read_bytes()).hexdigest(),
                   semantic_rules=[{"id": "fixture", "source_locator": "TEXT:Synthetic order", "summary": "fixture"}])
    source, output = tmp_path / "draft.docx", tmp_path / "ordered.docx"
    document = Document()
    for section in reversed(profile["sections"]):
        document.add_heading(section["label"], 1)
        document.add_paragraph("Payload " + section["id"])
        if section["id"] == "results":
            document.add_table(rows=1, cols=1).cell(0, 0).text = "Synthetic result table"
    document.save(source)
    result = apply_docx_structure(source, output, profile, author_approved=True)
    assert result["status"] == "STRUCTURE_APPLIED"
    assert result["validation"]["recognized"] == [section["id"] for section in profile["sections"]]
    generated = Document(output)
    assert generated.tables[0].cell(0, 0).text == "Synthetic result table"
    assert sorted(p.text for p in generated.paragraphs) == sorted(p.text for p in Document(source).paragraphs)


def test_vjpm_profile_limits_are_article_type_specific():
    profile = load_structure_profile(ROOT / "profiles/journal/vjpm-original-2026.yaml")
    assert profile["article_type"] == "original"
    assert profile["limits"]["total_words"]["enforcement"] == "recommendation"
    assert profile["limits"]["references"]["maximum"] == 18
    assert profile["sections"][-1]["id"] == "abstract-en"


def test_prisma_scr_includes_all_22_items_with_two_conditional_appraisals():
    manifest = load_coverage_manifest(ROOT / "coverage/prisma-scr-2018.yaml")
    assert manifest["expected_ids"] == [str(i) for i in range(1, 23)]
    assert manifest["implemented_ids"] == manifest["expected_ids"]
    assert [item["id"] for item in manifest["items"] if item.get("optional")] == ["12", "16"]
    assert manifest["official_tool_required_at_use"] is True
    assert validate_coverage_manifest(manifest) == []


@pytest.mark.parametrize("updates", [
    {"last_verified": "2999-01-01"}, {"expires_after_days": -1},
    {"expires_after_days": "invalid"}, {"last_verified": "invalid"},
])
def test_bad_freshness_metadata_never_verifies(tmp_path, updates):
    source = tmp_path / "source.txt"
    source.write_text("Synthetic rules", encoding="utf-8")
    profile = {"verification_status": "CURRENT", "last_verified": "2026-10-11",
               "expires_after_days": 180, "source_path": source.name,
               "source_sha256": sha256(source.read_bytes()).hexdigest(),
               "semantic_rules": [{"id": "fixture", "source_locator": "TEXT:Synthetic rules", "summary": "fixture"}]}
    profile.update(updates)
    assert verify_profile_source(tmp_path, profile)["status"] == "OFFICIAL_RULE_REQUIRED"


@pytest.mark.parametrize("filename,ids", [
    ("cross-2021.yaml", ["1a","1b","2","3","4","5a","5b","5c","5d","6a","6b","6c","6d","7a","7b","7c","8","9a","9b","10a","10b","10c","10d","10e","10f","10g","11a","11b","11c","11d","12","13a","13b","13c","14","15","16","17","18","19"]),
    ("record-2015.yaml", ["1.1","1.2","1.3","6.1","6.2","6.3","7.1","12.1","12.2","12.3","13.1","19.1","22.1"]),
    ("srqr-2014.yaml", [f"S{i}" for i in range(1,22)]),
    ("squire-2-2015.yaml", ["1","2a","2b","3","4","5","6","7","8a","8b","9a","9b","10a","10b","10c","11a","11b","12","13a","13b","13c","13d","13e","13f","14a","14b","15a","15b","15c","15d","15e","16a","16b","16c","17a","17b","17c","17d","17e","18"]),
])
def test_extensions_map_exact_primary_identifiers(filename, ids):
    manifest = load_coverage_manifest(ROOT / "coverage" / filename)
    assert manifest["expected_ids"] == ids
    assert manifest["implemented_ids"] == ids
    assert manifest["official_tool_required_at_use"] is True
    assert validate_coverage_manifest(manifest) == []


def test_cherries_uses_internal_ids_and_corrected_denominators():
    manifest = load_coverage_manifest(ROOT / 'coverage/cherries-2004-corrected-2012.yaml')
    assert manifest['expected_ids'] == [f'CHERRIES:{i:02}' for i in range(1, 31)]
    assert manifest['implemented_ids'] == manifest['expected_ids']
    assert manifest['identifiers_are_official'] is False
    assert manifest['rate_definitions'] == {
        'view': ['survey_visitors', 'site_visitors'],
        'participation': ['agreed', 'first_page_visitors'],
        'completion': ['finished', 'agreed'],
    }
    assert manifest['snapshot_status'] == 'LOCAL_DOWNLOAD_EMPTY'
    assert validate_coverage_manifest(manifest) == []
