import json
import zipfile
from itertools import count
from pathlib import Path
from types import SimpleNamespace

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree

from zotero_fields import (
    ZoteroFieldError,
    apply_edits,
    audit,
    build_citation_field,
    export_paragraphs,
    tokenize_paragraph,
)


INSTRUCTION = (
    'ADDIN ZOTERO_ITEM CSL_CITATION {"citationID":"a1b2c3","properties":{"formattedCitation":"(Nguyễn, 2020)",'
    '"plainCitation":"(Nguyễn, 2020)","noteIndex":0},"citationItems":[{"id":15,"uris":'
    '["http://zotero.org/users/123456/items/ABCD2345"],"itemData":{"id":15,"type":"article-journal",'
    '"title":"Example"}}],"schema":"https://github.com/citation-style-language/schema/raw/master/csl-citation.json"}'
)


def _run(paragraph, *children):
    run = OxmlElement("w:r")
    for child in children:
        run.append(child)
    paragraph._p.append(run)


def _fld(kind):
    node = OxmlElement("w:fldChar")
    node.set(qn("w:fldCharType"), kind)
    return node


def _instr(text):
    node = OxmlElement("w:instrText")
    node.set(qn("xml:space"), "preserve")
    node.text = text
    return node


def _t(text):
    node = OxmlElement("w:t")
    node.set(qn("xml:space"), "preserve")
    node.text = text
    return node


def add_field(paragraph, instruction, display, split=3):
    _run(paragraph, _fld("begin"))
    size = -(-len(instruction) // split)
    for start in range(0, len(instruction), size):
        _run(paragraph, _instr(instruction[start : start + size]))
    _run(paragraph, _fld("separate"))
    _run(paragraph, _t(display))
    _run(paragraph, _fld("end"))


def make_docx(path: Path, with_pref: bool = True) -> Path:
    document = Document()
    paragraph = document.add_paragraph("Tỷ lệ đáp ứng cao hơn ")
    add_field(paragraph, INSTRUCTION, "(Nguyễn, 2020)")
    paragraph.add_run(" so với nhóm chứng.")
    reference = document.add_paragraph("Xem ")
    add_field(reference, " REF _Ref123 \\h ", "Bảng 3.1", split=1)
    reference.add_run(" để biết chi tiết.")
    if with_pref:
        settings = document.settings.element
        variables = OxmlElement("w:docVars")
        variable = OxmlElement("w:docVar")
        variable.set(qn("w:name"), "ZOTERO_PREF_1")
        variable.set(qn("w:val"), '<data data-version="3"/>')
        variables.append(variable)
        settings.append(variables)
    document.save(path)
    return path


def _field_xml(path: Path, paragraph_index: int) -> list[bytes]:
    document = Document(str(path))
    paragraph = list(document.element.body.iter(qn("w:p")))[paragraph_index]
    pf, _ = tokenize_paragraph(paragraph, count(1))
    return [etree.tostring(run) for runs in pf.fields.values() for run in runs]


def test_round_trip_keeps_field_xml_byte_for_byte(tmp_path):
    source = make_docx(tmp_path / "source.docx")
    exported = export_paragraphs(source)
    first = exported["paragraphs"][0]
    assert first["text"] == "Tỷ lệ đáp ứng cao hơn ⟦Z:1⟧ so với nhóm chứng."
    assert exported["paragraphs"][1]["text"] == "Xem ⟦F:2⟧ để biết chi tiết."

    out = tmp_path / "out.docx"
    report = apply_edits(
        source,
        {first["index"]: "Tỷ lệ đáp ứng ở nhóm can thiệp cao hơn rõ rệt ⟦Z:1⟧ so với nhóm chứng."},
        out,
        expected_sha256=exported["source_sha256"],
    )

    assert report["status"] == "PASS"
    assert report["fields_before"] == report["fields_after"] == 1
    assert _field_xml(out, first["index"]) == _field_xml(source, first["index"])
    assert export_paragraphs(out)["paragraphs"][0]["text"].startswith("Tỷ lệ đáp ứng ở nhóm can thiệp cao hơn rõ rệt")


def test_deleted_token_blocks_the_edit_and_delivers_nothing(tmp_path):
    source = make_docx(tmp_path / "source.docx")
    out = tmp_path / "out.docx"
    with pytest.raises(ZoteroFieldError, match="FIELD_TOKEN_MISSING"):
        apply_edits(source, {0: "Tỷ lệ đáp ứng cao hơn so với nhóm chứng."}, out)
    assert not out.exists()


def test_cross_reference_fields_survive_too(tmp_path):
    source = make_docx(tmp_path / "source.docx")
    out = tmp_path / "out.docx"
    apply_edits(source, {1: "Chi tiết ở ⟦F:2⟧."}, out)
    assert _field_xml(out, 1) == _field_xml(source, 1)


def test_audit_detects_a_lost_field_and_a_lost_preference(tmp_path):
    source = make_docx(tmp_path / "source.docx")
    flattened = Document(str(source))
    flattened.paragraphs[0].text = "Tỷ lệ đáp ứng cao hơn (Nguyễn, 2020) so với nhóm chứng."
    for variables in flattened.settings.element.iter(qn("w:docVars")):
        variables.getparent().remove(variables)
    broken = tmp_path / "broken.docx"
    flattened.save(broken)

    report = audit(source, broken)

    assert report["status"] == "BLOCKED"
    assert len(report["missing"]) == 1
    assert report["missing"][0].startswith("ADDIN ZOTERO_ITEM CSL_CITATION")
    assert report["preferences_missing"] == ["ZOTERO_PREF_1"]


def test_zotero_guard_blocks_and_removes_a_file_that_lost_citations(tmp_path):
    from medical_research_skills_vn.docx_formatting import zotero_guard

    source = make_docx(tmp_path / "source.docx")
    flattened = Document(str(source))
    flattened.paragraphs[0].text = "Tỷ lệ đáp ứng cao hơn (Nguyễn, 2020) so với nhóm chứng."
    broken = tmp_path / "broken.docx"
    flattened.save(broken)

    report = zotero_guard(source, broken)

    assert report["status"] == "BLOCKED"
    assert not broken.exists()


def test_changed_source_is_refused(tmp_path):
    source = make_docx(tmp_path / "source.docx")
    with pytest.raises(ZoteroFieldError, match="SOURCE_CHANGED_SINCE_EXPORT"):
        apply_edits(source, {}, tmp_path / "out.docx", expected_sha256="0" * 64)


def test_paragraph_holding_an_image_is_not_editable(tmp_path):
    document = Document()
    paragraph = document.add_paragraph("Hình ")
    _run(paragraph, OxmlElement("w:drawing"))
    path = tmp_path / "image.docx"
    document.save(path)
    row = export_paragraphs(path)["paragraphs"][0]
    assert row["editable"] is False
    assert row["reason"] == "UNSUPPORTED_RUN_CONTENT"


def test_adjacent_ref_fields_sharing_a_boundary_run_are_rejected(tmp_path):
    document = Document()
    paragraph = document.add_paragraph("See ")
    _run(paragraph, _fld("begin"))
    _run(paragraph, _instr(" REF _RefOne \\h "))
    _run(paragraph, _fld("separate"))
    _run(paragraph, _t("Table 1"))
    _run(paragraph, _fld("end"), _fld("begin"))
    _run(paragraph, _instr(" REF _RefTwo \\h "))
    _run(paragraph, _fld("separate"))
    _run(paragraph, _t("Table 2"))
    _run(paragraph, _fld("end"))
    source = tmp_path / "adjacent.docx"
    out = tmp_path / "out.docx"
    document.save(source)

    row = export_paragraphs(source)["paragraphs"][0]
    assert row["editable"] is False
    with pytest.raises(ZoteroFieldError, match="PARAGRAPH_NOT_EDITABLE"):
        apply_edits(source, {0: row["text"] + " changed"}, out)
    assert not out.exists()


def test_audit_blocks_lost_ref_field(tmp_path):
    source = make_docx(tmp_path / "source.docx")
    flattened = Document(str(source))
    flattened.paragraphs[1].text = "Xem Bảng 3.1 để biết chi tiết."
    broken = tmp_path / "broken.docx"
    flattened.save(broken)

    report = audit(source, broken)
    assert report["status"] == "BLOCKED"
    assert report["fields_before"] == report["fields_after"] == 1
    assert any(instruction.startswith("REF _Ref123") for instruction in report["missing"])


def test_repeated_cite_placeholder_gets_distinct_live_fields(tmp_path):
    document = Document()
    document.add_paragraph("Cite twice.")
    source = tmp_path / "source.docx"
    out = tmp_path / "out.docx"
    document.save(source)
    item = SimpleNamespace(
        key="ABCD2345", item_id=15,
        uri="http://zotero.org/users/123456/items/ABCD2345",
        csl={"type": "article-journal", "author": [{"family": "Nguyễn"}], "issued": {"date-parts": [[2020]]}},
    )

    report = apply_edits(
        source,
        {0: "⟦cite:ABCD2345⟧ then ⟦cite:ABCD2345⟧"},
        out,
        resolver=lambda keys: ([item], []),
    )
    assert report["status"] == "PASS"
    assert report["added_count"] == 2
    with zipfile.ZipFile(out) as archive:
        root = etree.fromstring(archive.read("word/document.xml"))
    instructions = [node.text for node in root.iter(qn("w:instrText"))]
    payloads = [json.loads(text.split("CSL_CITATION ", 1)[1]) for text in instructions]
    assert len(payloads) == 2
    assert payloads[0]["citationID"] != payloads[1]["citationID"]


def test_corporate_and_incomplete_csl_authors_have_safe_labels():
    corporate = SimpleNamespace(
        key="CORP1234", item_id=1, uri="http://zotero.org/users/1/items/CORP1234",
        csl={"author": [{"literal": "World Health Organization"}], "issued": {"date-parts": [[2024]]}},
    )
    incomplete = SimpleNamespace(
        key="BARE1234", item_id=2, uri="http://zotero.org/users/1/items/BARE1234",
        csl={"author": [{}], "issued": {"date-parts": []}},
    )
    runs = build_citation_field([corporate, incomplete], "citation-1")
    display = "".join(node.text or "" for run in runs for node in run.iter(qn("w:t")))
    assert display == "(World Health Organization, 2024; BARE1234)"
